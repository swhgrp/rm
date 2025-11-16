#!/usr/bin/env python3
"""Generate blank document templates for OnlyOffice"""
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pathlib import Path

template_dir = Path("/opt/restaurant-system/files/src/files/templates/document_templates")
template_dir.mkdir(parents=True, exist_ok=True)

# Create blank Word document
doc = Document()
doc.save(template_dir / "blank.docx")
print("Created blank.docx")

# Create blank Excel spreadsheet
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
wb.save(template_dir / "blank.xlsx")
print("Created blank.xlsx")

# Create blank PowerPoint presentation
prs = Presentation()
prs.save(template_dir / "blank.pptx")
print("Created blank.pptx")

print(f"\nTemplates created in {template_dir}")
